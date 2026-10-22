"""Revised document text (plan section 8.3): the document with the ACCEPTED proposals applied.

* one accepted adjustment of a clause -> its ``proposed_text``;
* several accepted adjustments of the same clause -> one merge request (``merge-v1``) so the clause satisfies all;
* accepted ``add_section`` proposals are inserted after their anchor clause; ``new_document`` ones are listed only.
The result is a DRAFT (``.docx`` with changed / new paragraphs highlighted, and ``.md``) that the user edits in their
own editor and uploads as the next version. Nothing is published automatically.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from pydantic import BaseModel, field_validator

from kg.config import now_iso
from kg.llm import system_user
from kg.mapping.docview import DocView, load_doc
from kg.services.compare import word_diff
from kg.services.decisions import record_decision
from kg.services.permissions import require
from kg.steps import load_prompt, run_step

PROMPT = Path(__file__).parent.parent / "mapping" / "prompts" / "merge.md"
STEP = "merge"


class Contribution(BaseModel):
    proposal_id: int
    change: str = ""


class MergeItem(BaseModel):
    clause_id: str
    merged_text: str = ""
    conflicts: list[str] = []
    contributions: list[Contribution] = []

    @field_validator("conflicts", mode="before")
    @classmethod
    def _c(cls, v):
        return [] if v in (None, "") else ([v] if isinstance(v, str) else list(v))

    @field_validator("contributions", mode="before")
    @classmethod
    def _k(cls, v):
        return [] if v in (None, "") else [x for x in v if isinstance(x, dict) and x.get("proposal_id") is not None]


class MergeBatch(BaseModel):
    results: list[MergeItem] = []


def make_validator(ids: set[str]):
    def validate(out: MergeBatch) -> list[str]:
        out.results = [r for r in out.results if r.clause_id in ids]
        bad = [r.clause_id for r in out.results if not r.merged_text.strip()]
        miss = sorted(ids - {r.clause_id for r in out.results})
        return ([f"empty merged_text for {', '.join(bad)}"] if bad else []) + (
            ["missing results for clause ids: " + ", ".join(miss)] if miss else [])
    return validate


def depth(section_no: str | None) -> int:
    return min(4, (section_no or "1").count(".") + 1)


def merge_texts(store, cfg, run, llm, view: DocView, multi: dict[str, list[dict]], force: bool = False) -> dict:
    """One merge request per batch of clauses that several proposals change.
    -> {"text": {clause_id: merged}, "conflicts": {clause_id: [...]}, "notes": {clause_id: {proposal_id: str}},
        "requests": n}. A clause the model did not answer for is missing from "text"."""
    text_of, conflicts, notes, requests = {}, {}, {}, 0
    if not multi:
        return {"text": text_of, "conflicts": conflicts, "notes": notes, "requests": 0}
    system, version = load_prompt(PROMPT)
    size = int(cfg.get_in("mapping.merge_batch", cfg.get_in("mapping.proposal_batch", 5)))
    keys = sorted(multi)
    for i in range(0, len(keys), size):
        chunk = keys[i:i + size]
        payload = {"items": [{"clause_id": cid, "original": view.by_id[cid]["text"],
                              "proposals": [{"proposal_id": p["id"], "obligation": p.get("statement") or "",
                                             "text": p["proposed_text"], "explanation": p["explanation"]}
                                            for p in multi[cid]]}
                             for cid in chunk]}
        subject = f"{view.doc_id}:merge:" + hashlib.sha1("|".join(chunk).encode()).hexdigest()[:10]
        res = run_step(store, llm, STEP, subject, payload,
                       system_user(system, "INPUT:\n" + json.dumps(payload, ensure_ascii=False)), MergeBatch,
                       version, validate=make_validator(set(chunk)), force=force, run=run)
        requests += 1
        for r in (res.output.results if res.output else []):
            if not r.merged_text.strip():
                continue
            text_of[r.clause_id] = r.merged_text.strip()
            if r.conflicts:
                conflicts[r.clause_id] = r.conflicts
            ids = {p["id"] for p in multi[r.clause_id]}
            notes[r.clause_id] = {str(c.proposal_id): c.change.strip() for c in r.contributions
                                  if c.proposal_id in ids and c.change.strip()}
    return {"text": text_of, "conflicts": conflicts, "notes": notes, "requests": requests}


def build_blocks(store, cfg, run, llm, view: DocView, proposals: list[dict], force: bool = False,
                 edits: dict | None = None) -> dict:
    """-> {'blocks': [...], 'conflicts': {clause_id: [...]}, 'new_documents': [...], 'merge_requests': n}
    ``edits`` (``edits.get_edits``): a wording stored for exactly these proposals is used as it is."""
    from kg.services.edits import slot_of, valid_edit
    edits = edits or {}
    adjust: dict[str, list[dict]] = {}
    adds: dict[str, list[dict]] = {}
    new_docs = []

    def wording(p: dict) -> str:
        e = valid_edit(edits, slot_of(p), [p["id"]])
        return e["text"] if e else p["proposed_text"]

    for p in proposals:
        if p["kind"] == "adjust_clause" and p["clause_id"] in view.by_id:
            adjust.setdefault(p["clause_id"], []).append(p)
        elif p["kind"] == "add_section" and p["clause_id"] in view.by_id:
            adds.setdefault(p["clause_id"], []).append(p)
        elif p["kind"] == "new_document":
            new_docs.append({"heading": p["heading"], "text": wording(p), "obligation_id": p["obligation_id"]})
    text_of: dict[str, str] = {}
    conflicts: dict[str, list[str]] = {}
    multi = {}
    for cid, ps in adjust.items():
        e = valid_edit(edits, f"clause:{cid}", [p["id"] for p in ps])
        if e:
            text_of[cid] = e["text"]
        elif len(ps) == 1:
            text_of[cid] = ps[0]["proposed_text"]
        else:
            multi[cid] = ps
    merged = merge_texts(store, cfg, run, llm, view, multi, force)
    text_of.update(merged["text"])
    conflicts.update(merged["conflicts"])
    for cid in multi:                      # a failed merge: keep the first proposal, tell the user
        if cid not in text_of:
            text_of[cid] = multi[cid][0]["proposed_text"]
            conflicts[cid] = ["ავტომატური გაერთიანება ვერ მოხერხდა; გამოყენებულია პირველი წინადადება."]
    blocks = []
    for u in view.units:
        if u["unit_type"] == "signature":
            continue
        new = text_of.get(u["id"])
        blocks.append({"type": u["unit_type"], "id": u["id"], "section_no": u["section_no"],
                       "text": new if new is not None else u["text"], "original": u["text"],
                       "changed": new is not None and new.strip() != u["text"].strip(),
                       "proposal_ids": [p["id"] for p in adjust.get(u["id"], [])]})
        for p in adds.get(u["id"], []):
            blocks.append({"type": "added", "id": None, "section_no": None, "heading": p["heading"],
                           "text": wording(p), "original": "", "changed": True, "proposal_ids": [p["id"]]})
    return {"blocks": blocks, "conflicts": conflicts, "new_documents": new_docs, "merge_requests": merged["requests"]}


def active_proposals(store, version_id: str) -> list[dict]:
    """Proposals whose task is open or accepted (what the proposed version shows), with the obligation text."""
    rows = store.q("SELECT p.*, t.status AS task_status FROM proposals p JOIN tasks t "
                   "ON t.doc_version_id=p.doc_version_id AND t.obligation_id=p.obligation_id "
                   "WHERE p.doc_version_id=? AND t.status IN ('open','accepted') ORDER BY p.id", (version_id,))
    nodes = store.get_nodes([r["obligation_id"] for r in rows]) if rows else {}
    return [{**r, "statement": ((nodes.get(r["obligation_id"]) or {}).get("props") or {}).get("statement")}
            for r in rows]


def pending_merges(store, version_id: str) -> dict[str, list[dict]]:
    """Clauses that several active proposals change and that have no stored wording for exactly that set."""
    from kg.services.edits import get_edits, valid_edit
    edits = get_edits(store, version_id)
    by: dict[str, list[dict]] = {}
    for p in active_proposals(store, version_id):
        if p["kind"] == "adjust_clause" and p["clause_id"]:
            by.setdefault(p["clause_id"], []).append(p)
    return {cid: ps for cid, ps in by.items()
            if len(ps) > 1 and not valid_edit(edits, f"clause:{cid}", [p["id"] for p in ps])}


def merge_pending(cfg, store, run, llm, version_id: str, force: bool = False, only: set[str] | None = None) -> dict:
    """Merge (with the model) every clause that several active proposals change, and store the wording
    (``clause_edits``, source "merge"), so the proposed version shows ONE text per clause. A wording the user typed
    is never replaced. -> {"merged": n, "failed": n, "requests": n}"""
    from kg.services.edits import save_edit
    view = load_doc(store, version_id)
    multi = {cid: ps for cid, ps in pending_merges(store, version_id).items()
             if cid in view.by_id and (only is None or cid in only)}
    res = merge_texts(store, cfg, run, llm, view, multi, force)
    for cid, text in res["text"].items():
        notes = dict(res["notes"].get(cid, {}))
        if res["conflicts"].get(cid):
            notes["conflicts"] = res["conflicts"][cid]
        save_edit(store, None, version_id, f"clause:{cid}", text, [p["id"] for p in multi[cid]], "merge", notes)
    return {"merged": len(res["text"]), "failed": len(multi) - len(res["text"]), "requests": res["requests"]}


def to_markdown(view: DocView, built: dict) -> str:
    lines = [f"# {view.props.get('title') or view.doc_id}", ""]
    for b in built["blocks"]:
        if b["type"] == "metadata":
            continue
        if b["type"] == "heading":
            lines += ["", "#" * (depth(b["section_no"]) + 1) + " " + b["text"], ""]
        elif b["type"] == "added":
            lines += ["", f"**{b.get('heading') or 'ახალი პუნქტი'}** *(ახალი)*", "", b["text"], ""]
        else:
            prefix = f"{b['section_no']} " if b["type"] == "clause" and b["section_no"] else ""
            mark = " *(შეცვლილია)*" if b["changed"] else ""
            lines.append(f"{prefix}{b['text']}{mark}\n")
    if built["new_documents"]:
        lines += ["", "## ახალი დოკუმენტების წინადადებები", ""]
        lines += [f"- **{d['heading']}**: {d['text']}" for d in built["new_documents"]]
    return "\n".join(lines)


def to_docx(view: DocView, built: dict, path: Path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    doc = Document()
    doc.add_heading(view.props.get("title") or view.doc_id, 0)
    for b in built["blocks"]:
        if b["type"] == "metadata":
            continue
        if b["type"] == "heading":
            doc.add_heading(b["text"], depth(b["section_no"]))
            continue
        para = doc.add_paragraph()
        if b["type"] == "added":
            r = para.add_run((b.get("heading") or "ახალი პუნქტი") + ". ")
            r.bold = True
            r.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
            r2 = para.add_run(b["text"])
            r2.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
            continue
        prefix = f"{b['section_no']} " if b["type"] == "clause" and b["section_no"] else ""
        run = para.add_run(prefix + b["text"])
        if b["changed"]:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if built["new_documents"]:
        doc.add_heading("ახალი დოკუმენტების წინადადებები", 1)
        for d in built["new_documents"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run((d["heading"] or "") + ": ").bold = True
            p.add_run(d["text"])
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def changes_preview(built: dict) -> list[dict]:
    """[{id, original, new, diff}] for the changed clauses (shown before the download)."""
    return [{"id": b["id"], "original": b["original"], "new": b["text"], "diff": word_diff(b["original"], b["text"])}
            for b in built["blocks"] if b["type"] != "added" and b["changed"]]


def build_revised(cfg, store, run, llm, actor: dict, version_id: str, proposal_ids: list[int],
                  force: bool = False) -> dict:
    """Build the revised text from the given (accepted) proposals; writes ``.docx`` and ``.md``."""
    v = store.q1("SELECT dv.id, dv.family_id, d.company_id, d.code FROM document_versions dv "
                 "JOIN documents d ON d.id=dv.family_id WHERE dv.id=?", (version_id,))
    if v is None:
        raise KeyError(f"document version not found: {version_id}")
    require(actor, "decision.make", v["company_id"])
    view = load_doc(store, version_id)
    props = [p for p in store.q("SELECT * FROM proposals WHERE doc_version_id=? ORDER BY id", (version_id,))
             if p["id"] in set(proposal_ids)]
    from kg.services.edits import get_edits
    built = build_blocks(store, cfg, run, llm, view, props, force, edits=get_edits(store, version_id))
    stamp = now_iso().replace(":", "").replace("-", "")[:15]
    base = cfg.path("uploads") / v["company_id"] / "revised" / f"{v['code']}-{stamp}"
    to_docx(view, built, base.with_suffix(".docx"))
    base.with_suffix(".md").write_text(to_markdown(view, built), encoding="utf-8")
    rel = lambda p: str(p.relative_to(cfg.path("uploads")))  # noqa: E731
    out = {"docx": rel(base.with_suffix(".docx")), "md": rel(base.with_suffix(".md")),
           "changed_clauses": sum(1 for b in built["blocks"] if b["changed"] and b["type"] != "added"),
           "added": sum(1 for b in built["blocks"] if b["type"] == "added"),
           "conflicts": built["conflicts"], "new_documents": built["new_documents"],
           "merge_requests": built["merge_requests"], "changes": changes_preview(built)}
    record_decision(store, actor, "revise", doc_version_id=version_id,
                    data={k: out[k] for k in ("docx", "md", "changed_clauses", "added")} | {"proposals": proposal_ids})
    return out
